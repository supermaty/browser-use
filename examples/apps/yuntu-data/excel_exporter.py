"""
Excel导出模块 - 将整合后的数据导出为格式化的Excel文件

期望流程：
1. Agent 仅按提示词中的「取数路径」在云图页面上获取数据，不要求 Agent 输出特定文本格式。
2. 本模块根据「从 sample 文件分析得到的输出格式」将取到的数据填入模板，
   使用 docs/黑玩赛马复盘-template.xlsx 作为模板文件，完全遵照模板中的行列与样式。

当前支持两种模式：
- 模板模式（优先）：按 template_structure 分析结果固定行列写入，不依赖运行时搜索表头。
- 标准模式（回退）：完全由代码生成多 sheet 报表。若模板不存在则自动回退。
"""
from __future__ import annotations

from pathlib import Path
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from data_models import (
	YuntuDataReport,
	ProjectOverview,
	A5AssetFlow,
	KOLContentReview,
	SearchInsight,
	AdFlowReview,
	TAPortraitReview,
)
from utils import format_percentage

# ---------------------------------------------------------------------------
# 公共样式常量
# ---------------------------------------------------------------------------
_HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")

# ---------------------------------------------------------------------------
# 模板布局：来自 docs/【黑玩】赛马复盘0105.xlsx 的 template_structure 分析结果，
# 与 黑玩赛马复盘-template.xlsx 约定为同一结构，严格按此行/列填入，不搜索表头。
# ---------------------------------------------------------------------------
SHEET_OVERALL_NAME = "总体数据表现"
TEMPLATE_OVERALL_HEADER_ROW = 5   # 表头行（1-based）
TEMPLATE_OVERALL_FIRST_DATA_ROW = 6  # 第一条数据行（星图/TTL，1-based）

# (模板列 1-based, 来源 "meta"|"project", 模型字段名)
TEMPLATE_OVERALL_COLUMNS: list[tuple[int, str, str]] = [
	(3, "meta", "报告名称"),
	(4, "project", "消耗金额"),
	(5, "project", "曝光次数"),
	(6, "project", "曝光人数"),
	(7, "project", "互动率"),
	(8, "project", "互动量"),
	(9, "project", "A3流转人数"),
	(10, "project", "A3流转率"),
	(12, "project", "七日回搜人数"),
	(13, "project", "七日回搜率"),
	(14, "project", "CPM"),
	(15, "project", "CPE"),
	(16, "project", "CPA3"),
	(17, "project", "CP5A"),
	(18, "project", "CPS"),
	(20, "project", "完播率"),
]


# ---------------------------------------------------------------------------
# 公共辅助函数
# ---------------------------------------------------------------------------


def _apply_header_style(ws: Worksheet, row: int = 1) -> None:
	"""对指定行的所有已有单元格应用标准表头样式。"""
	for cell in ws[row]:
		cell.fill = _HEADER_FILL
		cell.font = _HEADER_FONT
		cell.alignment = _HEADER_ALIGNMENT


def _create_sheet_with_rows(
	wb: Workbook,
	sheet_name: str,
	headers: list[str],
	rows: list[tuple[str, Any]],
	col_widths: list[int] | None = None,
	position: int | None = None,
) -> Worksheet:
	"""
	创建含表头 + 数据行的 sheet（消除样板代码）。

	Args:
		wb: 目标 Workbook
		sheet_name: sheet 名称
		headers: 表头列名列表
		rows: [(label, value), ...] 数据行
		col_widths: 各列宽度（与 headers 等长），None 则使用默认宽度
		position: sheet 插入位置（None 则追加到末尾）
	"""
	ws = wb.create_sheet(sheet_name, position) if position is not None else wb.create_sheet(sheet_name)
	for col_idx, header in enumerate(headers, start=1):
		ws.cell(row=1, column=col_idx, value=header)
	_apply_header_style(ws)

	for row_idx, row_data in enumerate(rows, start=2):
		for col_idx, value in enumerate(row_data, start=1):
			cell = ws.cell(row=row_idx, column=col_idx, value=value if value is not None else "N/A")
			# 数值列右对齐
			if col_idx > 1:
				cell.alignment = Alignment(horizontal="right")

	if col_widths:
		from openpyxl.utils import get_column_letter
		for i, width in enumerate(col_widths, start=1):
			ws.column_dimensions[get_column_letter(i)].width = width

	return ws


# ---------------------------------------------------------------------------
# 入口函数
# ---------------------------------------------------------------------------


def export_to_excel(
	report: YuntuDataReport | list[YuntuDataReport],
	output_dir: Path | str,
	filename: str | None = None,
	template_path: Path | str | None = None,
) -> Path:
	"""
	将数据报告导出为格式化的Excel文件。
	支持单份报告或多份报告（黑玩多品牌分类：每分类星图+竞价），多份时模板模式按行依次填入「总体数据表现」。
	"""
	reports: list[YuntuDataReport] = [report] if isinstance(report, YuntuDataReport) else report
	output_dir = Path(output_dir)
	output_dir.mkdir(parents=True, exist_ok=True)

	if not filename:
		timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
		first = reports[0]
		brand = first.品牌名称 or "未知品牌"
		if len(reports) > 1:
			report_name = "多分类"
		else:
			report_name = first.报告名称 or "未知报告"
		filename = f"{brand}_{report_name}_{timestamp}.xlsx"
	if not filename.endswith('.xlsx'):
		filename += '.xlsx'
	file_path = output_dir / filename

	default_template = (
		Path(__file__).resolve().parent
		/ "docs"
		/ "黑玩赛马复盘-template.xlsx"
	)
	chosen_template: Path | None = None
	if template_path is not None:
		tp = Path(template_path)
		if tp.exists():
			chosen_template = tp
	elif default_template.exists():
		chosen_template = default_template

	if chosen_template is not None:
		try:
			_export_with_template(reports, chosen_template, file_path)
			return file_path
		except Exception:
			pass

	# ===== 标准模式 =====
	wb = Workbook()
	if 'Sheet' in wb.sheetnames:
		wb.remove(wb['Sheet'])
	# 多报告时：总体数据表现多行（一行一个报告），其余模块用第一份报告
	first = reports[0]
	if len(reports) > 1:
		_create_overview_multi_sheet(wb, reports)
	else:
		if first.项目整体:
			_create_project_overview_sheet(wb, first.项目整体)
	if first.五A人群资产流转:
		_create_a5_asset_flow_sheet(wb, first.五A人群资产流转)
	if first.达人及内容复盘:
		_create_kol_content_sheet(wb, first.达人及内容复盘)
	if first.搜索与溢出价值:
		_create_search_insight_sheet(wb, first.搜索与溢出价值)
	if first.投流数据精细化复盘:
		_create_ad_flow_sheet(wb, first.投流数据精细化复盘)
	if first.TA人群画像精准度复盘:
		_create_ta_portrait_sheet(wb, first.TA人群画像精准度复盘)
	_create_summary_sheet(wb, first)
	wb.save(file_path)
	return file_path


# ---------------------------------------------------------------------------
# 模板模式
# ---------------------------------------------------------------------------


def _export_with_template(
	reports: list[YuntuDataReport],
	template_path: Path,
	output_path: Path,
) -> None:
	"""
	按模板格式将多份报告填入：总体数据表现按行写入（第6行起每行一份报告），汇总取第一份。
	"""
	assert template_path.exists(), f"模板文件不存在: {template_path}"
	wb = load_workbook(template_path, data_only=False)
	if reports and SHEET_OVERALL_NAME in wb.sheetnames:
		ws = wb[SHEET_OVERALL_NAME]
		_fill_overview_multi_by_layout(ws, reports)
	if reports and "汇总" in wb.sheetnames:
		ws_summary = wb["汇总"]
		r0 = reports[0]
		try:
			ws_summary["B3"] = "、".join(r.报告名称 or "N/A" for r in reports) if len(reports) > 1 else (r0.报告名称 or "N/A")
			ws_summary["B4"] = r0.品牌名称 or "N/A"
			ws_summary["B5"] = r0.日期范围 or "N/A"
		except Exception:
			pass
	wb.save(output_path)


def _fill_overview_by_layout(ws: Worksheet, report: YuntuDataReport, data_row: int = TEMPLATE_OVERALL_FIRST_DATA_ROW) -> None:
	"""按固定行列将单份报告的 ProjectOverview 填入「总体数据表现」的指定行。"""
	project = report.项目整体
	if not project:
		return
	for col_1based, source, attr in TEMPLATE_OVERALL_COLUMNS:
		if source == "meta":
			value = getattr(report, attr, None)
		elif source == "project":
			value = getattr(project, attr, None)
		else:
			value = None
		if value is None:
			continue
		ws.cell(row=data_row, column=col_1based, value=value)


def _fill_overview_multi_by_layout(ws: Worksheet, reports: list[YuntuDataReport]) -> None:
	"""多份报告时，从 TEMPLATE_OVERALL_FIRST_DATA_ROW 起每行填一份报告。"""
	for i, report in enumerate(reports):
		_fill_overview_by_layout(ws, report, data_row=TEMPLATE_OVERALL_FIRST_DATA_ROW + i)


def _create_overview_multi_sheet(wb: Workbook, reports: list[YuntuDataReport]) -> None:
	"""标准模式多报告：创建「总体数据表现」sheet，表头行+多数据行（每行一份报告）。"""
	ws = wb.create_sheet("总体数据表现", 0)
	for col_1based, _source, attr in TEMPLATE_OVERALL_COLUMNS:
		ws.cell(row=TEMPLATE_OVERALL_HEADER_ROW, column=col_1based, value=attr)
	_apply_header_style(ws, row=TEMPLATE_OVERALL_HEADER_ROW)
	for i, report in enumerate(reports):
		_fill_overview_by_layout(ws, report, data_row=TEMPLATE_OVERALL_FIRST_DATA_ROW + i)


# ---------------------------------------------------------------------------
# 标准模式 - 各 Sheet 创建
# ---------------------------------------------------------------------------


def _create_project_overview_sheet(wb: Workbook, data: ProjectOverview) -> None:
	"""创建项目整体Overview sheet"""
	rows = [
		("消耗金额", data.消耗金额, "元"),
		("曝光次数", data.曝光次数, "次"),
		("曝光人数", data.曝光人数, "人"),
		("互动率", format_percentage(data.互动率) if data.互动率 else None, "%"),
		("互动量", data.互动量, "次"),
		("7日回搜人数", data.七日回搜人数, "人"),
		("七日回搜率", format_percentage(data.七日回搜率) if data.七日回搜率 else None, "%"),
		("回搜次数", data.回搜次数, "次"),
		("完播率", format_percentage(data.完播率) if data.完播率 else None, "%"),
		("完播数", data.完播数, "次"),
		("A3流转人数", data.A3流转人数, "人"),
		("A3流转率", format_percentage(data.A3流转率) if data.A3流转率 else None, "%"),
		("本次活动转化金额", data.本次活动转化金额, "元"),
		("拉新人群规模", data.拉新人群规模, "人"),
		("新客占比", format_percentage(data.新客占比) if data.新客占比 else None, "%"),
		("CPM", data.CPM, "元/千次曝光"),
		("CPE", data.CPE, "元/次互动"),
		("CPS", data.CPS, "元/次回搜"),
		("CPA3", data.CPA3, "元/A3流转"),
		("CP5A", data.CP5A, "元/5A人群"),
		("ROI", data.ROI, "倍"),
	]
	_create_sheet_with_rows(
		wb, "项目整体", ["指标", "数值", "单位/说明"],
		rows, col_widths=[20, 20, 15], position=0,
	)


def _create_a5_asset_flow_sheet(wb: Workbook, data: A5AssetFlow) -> None:
	"""创建5A人群资产流转 sheet"""
	rows = [
		("投后人群规模", data.投后人群规模),
		("投后增长率", format_percentage(data.投后增长率) if data.投后增长率 else None),
		("A1新增量级", data.A1新增量级),
		("A1流转率", format_percentage(data.A1流转率) if data.A1流转率 else None),
		("A2新增量级", data.A2新增量级),
		("A2流转率", format_percentage(data.A2流转率) if data.A2流转率 else None),
		("A3新增量级", data.A3新增量级),
		("A3流转率", format_percentage(data.A3流转率) if data.A3流转率 else None),
		("A4新增量级", data.A4新增量级),
		("A4流转率", format_percentage(data.A4流转率) if data.A4流转率 else None),
		("A5新增量级", data.A5新增量级),
		("A5流转率", format_percentage(data.A5流转率) if data.A5流转率 else None),
	]
	_create_sheet_with_rows(
		wb, "5A人群资产流转", ["指标", "数值"],
		rows, col_widths=[20, 20],
	)


def _create_kol_content_sheet(wb: Workbook, data: KOLContentReview) -> None:
	"""创建达人及内容复盘 sheet"""
	rows = [
		("Excel文件路径", data.excel_file_path or "未下载"),
		("是否已下载", "是" if data.excel_downloaded else "否"),
		("内容数量", data.内容数量),
		("爆文率", format_percentage(data.爆文率) if data.爆文率 else None),
		("看后搜索率", format_percentage(data.看后搜索率) if data.看后搜索率 else None),
	]
	_create_sheet_with_rows(
		wb, "达人及内容复盘", ["项目", "内容"],
		rows, col_widths=[20, 50],
	)


def _create_search_insight_sheet(wb: Workbook, data: SearchInsight) -> None:
	"""创建搜索与溢出价值 sheet"""
	rows = [
		("SOV（声量份额）", data.SOV),
		("搜索人数变化趋势", format_percentage(data.搜索人数变化趋势) if data.搜索人数变化趋势 else None),
		("搜索次数变化趋势", format_percentage(data.搜索次数变化趋势) if data.搜索次数变化趋势 else None),
	]
	_create_sheet_with_rows(
		wb, "搜索与溢出价值", ["指标", "数值"],
		rows, col_widths=[25, 20],
	)


def _create_ad_flow_sheet(wb: Workbook, data: AdFlowReview) -> None:
	"""创建投流数据精细化复盘 sheet"""
	rows = [
		("Excel文件路径", data.excel_file_path or "未下载"),
		("是否已下载", "是" if data.excel_downloaded else "否"),
		("投放A3流转率", format_percentage(data.投放A3流转率) if data.投放A3流转率 else None),
		("人群精准度", format_percentage(data.人群精准度) if data.人群精准度 else None),
	]
	_create_sheet_with_rows(
		wb, "投流数据精细化复盘", ["项目", "内容"],
		rows, col_widths=[20, 50],
	)


def _create_ta_portrait_sheet(wb: Workbook, data: TAPortraitReview) -> None:
	"""创建TA人群画像精准度复盘 sheet"""
	ws = wb.create_sheet("TA人群画像")
	for col_idx, header in enumerate(["人群类型", "画像类型", "内容"], start=1):
		ws.cell(row=1, column=col_idx, value=header)
	_apply_header_style(ws)

	rows = [
		("全触达人群", "基础画像", str(data.全触达人群基础画像) if data.全触达人群基础画像 else "N/A"),
		("全触达人群", "内容偏好", str(data.全触达人群内容偏好) if data.全触达人群内容偏好 else "N/A"),
		("投后A3人群", "基础画像", str(data.投后A3人群基础画像) if data.投后A3人群基础画像 else "N/A"),
		("投后A3人群", "内容偏好", str(data.投后A3人群内容偏好) if data.投后A3人群内容偏好 else "N/A"),
	]

	for idx, (crowd, portrait_type, content) in enumerate(rows, start=2):
		ws.cell(row=idx, column=1, value=crowd)
		ws.cell(row=idx, column=2, value=portrait_type)
		cell = ws.cell(row=idx, column=3, value=content)
		cell.alignment = Alignment(wrap_text=True, vertical="top")

	from openpyxl.utils import get_column_letter
	ws.column_dimensions['A'].width = 15
	ws.column_dimensions['B'].width = 15
	ws.column_dimensions['C'].width = 60


def _create_summary_sheet(wb: Workbook, report: YuntuDataReport) -> None:
	"""创建汇总sheet"""
	ws = wb.create_sheet("汇总", 0)

	# 标题
	ws['A1'] = "数据报告汇总"
	ws['A1'].font = Font(bold=True, size=16)

	# 基本信息
	ws['A3'] = "报告名称："
	ws['B3'] = report.报告名称 or "N/A"
	ws['A4'] = "品牌名称："
	ws['B4'] = report.品牌名称 or "N/A"
	ws['A5'] = "日期范围："
	ws['B5'] = report.日期范围 or "N/A"

	# 数据模块状态
	ws['A7'] = "数据模块"
	ws['B7'] = "状态"
	_apply_header_style(ws, row=7)

	modules = [
		("项目整体", report.项目整体 is not None),
		("5A人群资产流转", report.五A人群资产流转 is not None),
		("达人及内容复盘", report.达人及内容复盘 is not None),
		("搜索与溢出价值", report.搜索与溢出价值 is not None),
		("投流数据精细化复盘", report.投流数据精细化复盘 is not None),
		("TA人群画像精准度复盘", report.TA人群画像精准度复盘 is not None),
	]

	for idx, (module_name, has_data) in enumerate(modules, start=8):
		ws[f'A{idx}'] = module_name
		ws[f'B{idx}'] = "✓ 已提取" if has_data else "✗ 未提取"

	ws.column_dimensions['A'].width = 25
	ws.column_dimensions['B'].width = 15
